"""Name matching for the pricing-tracker import, and prior-file replacement."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd

from src import history_archive
from src.branding import search_label
from src.pricing_tracker import (
    aliases_for,
    build_observed_history,
    import_observed_brand,
    match_tracker_name,
    normalise_dish_name,
    read_big_table,
    row_brand_for,
)


SCRAPE_NAMES = [
    "10 Chicken Wings",
    "5 Chicken Wings",
    "PERi-PERi Nuts",
    "Butterfly Burger",
    "Bottle of Coke Zero",
    "Spicy Rice",
    "One Pastel de Nata",
]


def test_normalise_treats_peri_spellings_as_the_same_dish() -> None:
    assert normalise_dish_name("PERi-PERi Nuts") == normalise_dish_name("Peri-Peri Nuts")
    assert normalise_dish_name("Tenderstem® Broccoli") == normalise_dish_name(
        "Tenderstem Broccoli"
    )
    assert normalise_dish_name("Ham & Mozzarella") == normalise_dish_name(
        "Ham and Mozzarella"
    )
    assert normalise_dish_name("Merlot (250ml)") == normalise_dish_name("Merlot 250ml")
    assert normalise_dish_name("chocolate brownie (v)") == normalise_dish_name(
        "chocolate brownie"
    )
    assert normalise_dish_name("Tall Husband's favourite") == normalise_dish_name(
        "Tall husbands favourite"
    )


def test_match_refuses_a_different_dish_with_a_similar_name() -> None:
    aliases = aliases_for("Nando's")

    assert match_tracker_name("15 Chicken Wings", SCRAPE_NAMES, aliases) is None
    assert match_tracker_name("Bottomless Coke Zero", SCRAPE_NAMES, aliases) is None
    assert (
        match_tracker_name("Spicy Rice (Northern Ireland)", SCRAPE_NAMES, aliases) is None
    )
    assert match_tracker_name("Peri-Peri Nuts", SCRAPE_NAMES, aliases) == "PERi-PERi Nuts"
    assert match_tracker_name("Buttefly Burger", SCRAPE_NAMES, aliases) == "Butterfly Burger"


def test_build_merges_duplicate_rows_and_rejects_conflicting_prices() -> None:
    menu = pd.DataFrame(
        {"name": ["PERi-PERi Nuts"], "category": ["Starters"], "price": [4.25]}
    )
    prices = pd.DataFrame(
        {
            "tracker_name": ["Peri-Peri Nuts", "PERi-PERi Nuts"],
            "date": [date(2025, 3, 1), date(2025, 3, 1)],
            "price": [3.95, 3.95],
        }
    )

    report = build_observed_history(prices, menu, "Nando's", aliases={})

    assert len(report.history) == 1
    assert report.history.loc[0, "item"] == "PERi-PERi Nuts"
    assert report.history.loc[0, "price"] == 3.95

    conflict = prices.assign(price=[3.95, 4.25])
    clashed = build_observed_history(conflict, menu, "Nando's", aliases={})
    assert clashed.history.loc[0, "price"] == 4.25
    assert clashed.conflicts == (("PERi-PERi Nuts", "2025-03-01", 4.25, 3.95),)


def test_read_big_table_ignores_summary_rows_and_other_brands(tmp_path: Path) -> None:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Nandos"
    sheet["B1"] = "Nandos"
    sheet["B3"] = "Average"
    sheet.append(["Brand", "Course", "Combo", "Dish name", date(2022, 2, 1), "vs Feb-22"])
    sheet.append(["Nando's", "Mains", "Mains", "1/2 Chicken", 7.95, 0.16])
    sheet.append(["Pizza Express", "Mains", "Mains", "Margherita", 11.0, 0.1])
    path = tmp_path / "tracker.xlsx"
    book.save(path)

    table = read_big_table(path, "Nandos", "Nando's")

    assert table["tracker_name"].tolist() == ["1/2 Chicken"]
    assert table.loc[0, "price"] == 7.95
    assert table.loc[0, "date"] == date(2022, 2, 1)


def test_rosas_sheet_uses_the_tracker_brand_not_the_stored_one(tmp_path: Path) -> None:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Rosas"
    sheet.append(["Brand", "Course", "Combo", "Dish name", date(2025, 3, 1)])
    sheet.append(["Rosas", "Mains", "Mains", "Sticky Rice", 4.1])
    sheet.append(["Rosa's Thai", "Mains", "Mains", "Other Dish", 9.0])
    path = tmp_path / "tracker.xlsx"
    book.save(path)

    assert row_brand_for("Rosa's Thai") == "Rosas"
    table = read_big_table(path, "Rosas", "Rosa's Thai", row_brand="Rosas")

    assert table["tracker_name"].tolist() == ["Sticky Rice"]


def test_import_observed_brand_writes_prior_and_refuses_an_empty_match(
    tmp_path: Path,
) -> None:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Rosas"
    sheet.append(["Brand", "Course", "Combo", "Dish name", date(2025, 3, 1)])
    sheet.append(["Rosas", "Mains", "Mains", "Sticky Rice", 4.1])
    path = tmp_path / "tracker.xlsx"
    book.save(path)
    prior = tmp_path / "price_history.csv"
    menu = pd.DataFrame({"name": ["Sticky Rice"], "category": ["Rice"], "price": [4.4]})

    written = import_observed_brand(path, "Rosa's Thai", menu, prior_path=prior)

    assert written["written"] is True
    assert written["dishes"] == 1
    saved = pd.read_csv(prior)
    assert saved["source"].tolist() == ["Rosa's Thai"]
    assert saved["item"].tolist() == ["Sticky Rice"]

    empty_menu = pd.DataFrame({"name": ["Not on the tracker"], "category": ["Mains"], "price": [1.0]})
    skipped = import_observed_brand(path, "Rosa's Thai", empty_menu, prior_path=prior)
    assert skipped["written"] is False
    assert pd.read_csv(prior)["item"].tolist() == ["Sticky Rice"]


def test_search_label_keeps_the_real_name_and_adds_the_plain_spelling() -> None:
    assert search_label("Banana Tree") == "Banana Tree"
    assert search_label("Rosa's Thai") == "Rosa's Thai · Rosas Thai"
    assert "Rosas" in search_label("Rosa's Thai (London - Islington, Restaurant)")


def test_replace_brand_in_prior_leaves_other_brands(tmp_path: Path) -> None:
    prior = tmp_path / "price_history.csv"
    existing = pd.DataFrame(
        {
            "date": [date(2022, 3, 31), date(2022, 3, 31)],
            "source": ["Nando's", "Pizza Express"],
            "item": ["Chips", "Margherita"],
            "dish_id": ["old", "px"],
            "bucket": ["Sides", "Mains"],
            "price": [1.0, 11.0],
        }
    )
    existing.to_csv(prior, index=False)
    incoming = pd.DataFrame(
        {
            "date": [date(2022, 2, 1)],
            "source": ["Nando's"],
            "item": ["Chips"],
            "dish_id": ["new"],
            "bucket": ["Sides"],
            "price": [3.5],
        }
    )

    counts = history_archive.replace_brand_in_prior("Nando's", incoming, prior)

    saved = pd.read_csv(prior)
    assert counts["removed"] == 1
    assert counts["kept"] == 1
    assert saved.loc[saved["source"] == "Pizza Express", "item"].tolist() == ["Margherita"]
    assert saved.loc[saved["source"] == "Nando's", "price"].tolist() == [3.5]
    assert history_archive.load_observed_sources(prior) == {"Nando's"}
